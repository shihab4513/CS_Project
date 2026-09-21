"""Create a presentation update for the integrated label-efficiency experiment."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "04_documents" / "Revised_AI_Based_Malicious_Package_Detection_Project_Presentation.pptx"
OUTPUT = ROOT / "04_documents" / "Revised_AI_Based_Malicious_Package_Detection_Project_Presentation_v2.pptx"


def add_slide(deck: Presentation, title: str, subtitle: str, sections: list[tuple[str, str]]) -> None:
    blank_layout = next(
        (layout for layout in deck.slide_layouts if "blank" in layout.name.lower()),
        deck.slide_layouts[-1],
    )
    slide = deck.slides.add_slide(blank_layout)
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = RGBColor(247, 249, 252)

    title_box = slide.shapes.add_textbox(Inches(0.65), Inches(0.42), Inches(12.0), Inches(0.6))
    title_frame = title_box.text_frame
    title_frame.text = title
    title_frame.paragraphs[0].font.size = Pt(28)
    title_frame.paragraphs[0].font.bold = True
    title_frame.paragraphs[0].font.color.rgb = RGBColor(16, 35, 63)

    subtitle_box = slide.shapes.add_textbox(Inches(0.68), Inches(1.05), Inches(11.8), Inches(0.45))
    subtitle_frame = subtitle_box.text_frame
    subtitle_frame.text = subtitle
    subtitle_frame.paragraphs[0].font.size = Pt(13)
    subtitle_frame.paragraphs[0].font.color.rgb = RGBColor(75, 91, 111)

    top = 1.75
    for heading, body in sections:
        box = slide.shapes.add_textbox(Inches(0.85), Inches(top), Inches(11.4), Inches(0.9))
        frame = box.text_frame
        frame.word_wrap = True
        paragraph = frame.paragraphs[0]
        paragraph.text = heading
        paragraph.font.size = Pt(17)
        paragraph.font.bold = True
        paragraph.font.color.rgb = RGBColor(28, 114, 147)
        detail = frame.add_paragraph()
        detail.text = body
        detail.font.size = Pt(15)
        detail.font.color.rgb = RGBColor(38, 49, 64)
        detail.space_after = Pt(8)
        top += 1.25

    footer = slide.shapes.add_textbox(Inches(0.7), Inches(7.05), Inches(12), Inches(0.25))
    footer.text_frame.text = "AI-Based Malicious Open-Source Package Detection | Integrated Experiment C"
    footer.text_frame.paragraphs[0].font.size = Pt(9)
    footer.text_frame.paragraphs[0].font.color.rgb = RGBColor(110, 123, 140)
    footer.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT


def main() -> None:
    deck = Presentation(SOURCE)
    add_slide(
        deck,
        "Updated XGBoost Transfer Result",
        "The latest executable run reports the model and direction explicitly.",
        [
            ("NPM to PyPI", "F1 = 0.4662 | ROC-AUC = 0.7277"),
            ("PyPI to NPM", "F1 = 0.3058 | ROC-AUC = 0.7887"),
            ("Interpretation", "Transfer remains asymmetric. These values are XGBoost results from the current run and are not deployment estimates."),
        ],
    )
    add_slide(
        deck,
        "Experiment C: Target-Label Efficiency",
        "How many labeled packages from a new ecosystem are needed?",
        [
            ("Condition A", "Train on all source-ecosystem packages plus k labeled target packages."),
            ("Condition B", "Train on k target packages only. Both conditions use the same held-out target set."),
            ("Protocol", "8 seeds, 65% target pool, 35% target holdout, budgets from 0 to 400, with F1, ROC-AUC, and PR-AUC."),
        ],
    )
    add_slide(
        deck,
        "Experiment C: Initial Verified Results",
        "NPM source data remains useful, but target labels progressively reduce the gap.",
        [
            ("NPM plus k PyPI labels", "F1 = 0.533 at k=0; 0.586 at k=25; 0.652 at k=50; 0.696 at k=200; 0.711 at k=400."),
            ("Target-only comparison", "At k=400, PyPI-only F1 = 0.698 versus 0.711 with NPM data added."),
            ("Scientific reading", "The source corpus gives a useful transfer floor, but this experiment does not prove that 25 labels are sufficient."),
        ],
    )
    add_slide(
        deck,
        "Final Contribution After Integration",
        "The contribution is an evaluation and decision-policy framework, not a new hybrid detector.",
        [
            ("Contribution", "Measure directional transfer, target-label efficiency, dynamic-analysis referral, modality disagreement, and evidence provenance."),
            ("Operational value", "Show when source data helps and when a registry should invest in target-ecosystem labeling."),
            ("Limitations", "One released dataset, limited dynamic coverage, small triage example, and no claim that LLM output is ground truth."),
        ],
    )
    deck.core_properties.title = "AI-Based Malicious Open-Source Package Detection - Integrated Results"
    deck.core_properties.subject = "Cross-ecosystem transfer and target-label efficiency"
    deck.save(OUTPUT)
    print(f"Wrote updated presentation to: {OUTPUT}")


if __name__ == "__main__":
    main()